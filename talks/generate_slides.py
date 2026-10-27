"""Generate a GroupBees deck (.pptx) from slides_content.py.

Usage (from the repo root):
    uv run --with python-pptx python talks/generate_slides.py

Output: talks/slides/<DECK_NAME>.pptx — DECK_NAME is set in slides_content.py.

Design: GroupBees identity — cyan → purple gradient, near-black ink, white.
Assets in talks/assets/ (logo + bee mascots). Slide types: title, speakers,
journey, bullets, compare, image, code, demo, takeaways, end.

Every slide also takes `bees`: decorative mascots, as tuples
(file, x, y, height[, rotation_deg[, flip]]) in inches — add your own freely.

Motion: every slide gets the Gallery transition (1 s) and every bee a Zoom
entrance that starts with the slide ("With Previous"). Neither the Google
Slides API nor Apps Script can set these, so they ship in the .pptx; check them
after importing into Google Slides. Set TRANSITION / BEE_ANIMATION to False in
slides_content.py to turn them off.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

import slides_content
from slides_content import FOOTER_TEXT, SLIDES

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "slides" / f"{getattr(slides_content, 'DECK_NAME', 'deck')}.pptx"

# ── GroupBees palette (sampled from groupbees.fr) ────────────────────────────
INK = RGBColor(0x14, 0x14, 0x1A)        # near-black of the wordmark
CYAN = RGBColor(0x00, 0xB5, 0xFB)       # gradient start
PURPLE = RGBColor(0x97, 0x4E, 0xFF)     # gradient end
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xC9, 0xCF, 0xDC)       # muted text on dark
GRAY = RGBColor(0x5B, 0x63, 0x72)       # muted text on light
BODY_INK = RGBColor(0x2A, 0x2F, 0x3A)   # bullet / body text
CARD = RGBColor(0x1B, 0x1B, 0x24)       # code card
CODE_FG = RGBColor(0xE8, 0xEA, 0xF2)
HEX_TINT = RGBColor(0xEB, 0xF6, 0xFE)   # honeycomb watermark

HEADING_FONT = "Avenir Next"
BODY_FONT = "Avenir Next"
CODE_FONT = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── low-level helpers ────────────────────────────────────────────────────────
def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _set_run(run, *, size=None, bold=None, italic=None, color=None, font=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = font or BODY_FONT


def _textbox(slide, x, y, w, h, *, anchor=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tf


def _line(tf, text, *, first=False, align=PP_ALIGN.LEFT, space_after=0, **run_kw):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run(run, **run_kw)
    return p


def _rect(slide, x, y, w, h, color, *, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _gradient(slide, x, y, w, h, *, angle=0.0, shape=MSO_SHAPE.RECTANGLE):
    """Cyan → purple gradient shape — the GroupBees signature accent."""
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.gradient()
    s.fill.gradient_angle = angle
    stops = s.fill.gradient_stops
    stops[0].color.rgb = CYAN
    stops[0].position = 0.0
    stops[1].color.rgb = PURPLE
    stops[1].position = 1.0
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _dark_background(slide):
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, INK)


def _honeycomb(slide, x, y, *, cols=2, rows=3, size=Inches(0.62), color=HEX_TINT):
    """Light honeycomb watermark (a nod to the bees, kept very subtle)."""
    dx = size * 0.86
    dy = size * 0.75
    for r in range(rows):
        for c in range(cols):
            _rect(
                slide,
                int(x + c * dx + (r % 2) * dx / 2),
                int(y + r * dy),
                int(size),
                int(size),
                color,
                shape=MSO_SHAPE.HEXAGON,
            )


def _picture(slide, name, *, left=None, top=None, width=None, height=None):
    path = ASSETS / name
    if not path.exists():
        return None
    pic = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    pic.name = name  # bee*.png pictures get the Zoom entrance (see _animate_bees)
    return pic


def _bees(slide, data):
    """Decorative bees: (file, x, y, h[, rotation[, flip]]) in inches."""
    for spec in data.get("bees", []):
        name, x, y, h, *rest = spec
        rot = rest[0] if rest else 0
        flip = rest[1] if len(rest) > 1 else False
        pic = _picture(slide, name, left=Inches(x), top=Inches(y), height=Inches(h))
        if pic is None:
            continue
        pic.rotation = rot
        if flip:
            pic._element.spPr.find(qn("a:xfrm")).set("flipH", "1")


def _footer(slide, index, total, *, dark=False):
    """Logo left, talk title + slide number right."""
    if index == 0:
        return
    _picture(slide, "logo_light.png" if dark else "logo.png",
             left=Inches(0.62), top=Inches(6.85), height=Inches(0.3))
    tf = _textbox(slide, Inches(7.5), Inches(6.88), Inches(5.2), Inches(0.35))
    _line(tf, f"{FOOTER_TEXT}   ·   {index}/{total}", first=True,
          align=PP_ALIGN.RIGHT, size=11, color=MIST if dark else GRAY)


def _title(slide, text, *, dark=False):
    tf = _textbox(slide, Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.0))
    _line(tf, text, first=True, size=34, bold=True,
          color=WHITE if dark else INK, font=HEADING_FONT)
    _gradient(slide, Inches(0.7), Inches(1.32), Inches(1.5), Inches(0.08))


def _notes(slide, notes):
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _cue(slide, cue):
    """Put the stage cue (who speaks, hand-over) above the slide's notes."""
    tf = slide.notes_slide.notes_text_frame
    tf.text = f"{cue}\n\n────────\n{tf.text}" if tf.text else cue


def _bullet_hex(slide, y, i):
    """Small alternating cyan/purple hexagon used as a bullet marker."""
    color = CYAN if i % 2 == 0 else PURPLE
    _rect(slide, Inches(0.78), y, Inches(0.19), Inches(0.19), color,
          shape=MSO_SHAPE.HEXAGON)


# ── renderers ────────────────────────────────────────────────────────────────
def render_title(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    _picture(slide, "logo_light.png", left=Inches(0.7), top=Inches(0.62),
             width=Inches(2.5))
    _picture(slide, data.get("mascot", "bee.png"), left=Inches(9.4),
             top=Inches(1.5), height=Inches(4.6))

    tf = _textbox(slide, Inches(0.7), Inches(1.9), Inches(8.6), Inches(2.0),
                  anchor=MSO_ANCHOR.BOTTOM)
    for k, part in enumerate(data["title"].split("\n")):
        _line(tf, part, first=(k == 0), size=data.get("title_size", 54),
              bold=True, color=WHITE, font=HEADING_FONT)
    if data.get("subtitle"):
        tf2 = _textbox(slide, Inches(0.7), Inches(4.1), Inches(8.2), Inches(1.0))
        _line(tf2, data["subtitle"], first=True, size=24, color=MIST)
    _gradient(slide, Inches(0.7), Inches(5.15), Inches(2.4), Inches(0.1))
    if data.get("footer"):
        tf3 = _textbox(slide, Inches(0.7), Inches(5.6), Inches(8.6), Inches(0.5))
        _line(tf3, data["footer"], first=True, size=16, color=CYAN)
    _notes(slide, data.get("notes"))


def render_end(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    qrs = getattr(slides_content, "END_QRS", [])
    if qrs:
        _end_qrs(slide, qrs)
    else:
        _picture(slide, data.get("mascot", "bee.png"), left=Inches(9.6),
                 top=Inches(1.8), height=Inches(4.2))

    tf = _textbox(slide, Inches(0.7), Inches(2.4), Inches(8.6), Inches(1.4))
    _line(tf, data["title"], first=True, size=54, bold=True, color=WHITE,
          font=HEADING_FONT)
    _gradient(slide, Inches(0.7), Inches(3.9), Inches(2.4), Inches(0.1))
    tf2 = _textbox(slide, Inches(0.7), Inches(4.3), Inches(8.2), Inches(2.0))
    for i, line in enumerate(data.get("lines", [])):
        _line(tf2, line, first=(i == 0), size=20, color=MIST, space_after=10)
    _notes(slide, data.get("notes"))


def _end_qrs(slide, qrs):
    """QR codes on the end slide, in place of the mascot: the first one large
    (the event's feedback form), the others small in a row beneath it.
    Each is (file, caption); a white card keeps the quiet zone scannable."""
    def qr(name, caption, x, y, size, caption_size):
        pad = Inches(0.12)
        _rect(slide, x - pad, y - pad, size + 2 * pad, size + 2 * pad, WHITE,
              shape=MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.05
        _picture(slide, name, left=x, top=y, width=size, height=size)
        tf = _textbox(slide, x - Inches(0.5), y + size + Inches(0.18),
                      size + Inches(1.0), Inches(0.4))
        _line(tf, caption, first=True, align=PP_ALIGN.CENTER, size=caption_size,
              color=MIST)

    (name, caption), *others = qrs
    big = Inches(3.3)
    qr(name, caption, Inches(9.45), Inches(1.2), big, 16)
    small, gap = Inches(1.2), Inches(0.5)
    x = Inches(9.45) + (big - len(others) * small - (len(others) - 1) * gap) // 2
    for name, caption in others:
        qr(name, caption, x, Inches(5.35), small, 11)
        x += small + gap


def render_bullets(prs, data):
    slide = _blank(prs)
    _honeycomb(slide, Inches(11.45), Inches(1.9))
    _title(slide, data["title"])

    image = data.get("image")
    text_w = Inches(7.7) if image else Inches(11.5)
    if image:
        _picture(slide, image, left=Inches(9.6), top=Inches(2.2),
                 height=Inches(3.2) if image != "usbc.png" else None,
                 width=Inches(3.2) if image == "usbc.png" else None)

    # ~ chars per line at 21pt in the available width; used to reserve room for
    # wrapped bullets so the list never bunches up.
    per_line = int(text_w / Inches(0.132))
    y = Inches(1.95)
    for i, bullet in enumerate(data["bullets"]):
        _bullet_hex(slide, y + Inches(0.12), i)
        tf = _textbox(slide, Inches(1.2), y, text_w, Inches(0.6))
        _line(tf, bullet, first=True, size=21, color=BODY_INK)
        lines = max(1, -(-len(bullet) // per_line))
        y += Inches(0.70) + Inches(0.38) * (lines - 1)

    if data.get("kicker"):
        _gradient(slide, Inches(0.7), Inches(6.2), Inches(11.9), Inches(0.05))
        tf = _textbox(slide, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5))
        _line(tf, data["kicker"], first=True, size=18, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


def render_image(prs, data):
    slide = _blank(prs)
    _title(slide, data["title"])
    path = ASSETS.parent / data["image_path"] if data.get("image_path") else None
    if path and not path.is_absolute():
        path = (ROOT / data["image_path"]).resolve()

    body_top, body_h, body_w = Inches(1.65), Inches(4.9), Inches(11.9)
    if path and path.exists():
        pic = slide.shapes.add_picture(str(path), Inches(0.7), body_top, height=body_h)
        if pic.width > body_w:
            ratio = body_w / pic.width
            pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
        pic.left = int((SLIDE_W - pic.width) / 2)
    else:
        box = _rect(slide, Inches(0.7), body_top, body_w, body_h, HEX_TINT)
        tf = box.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _line(tf, f"[ {data.get('image_placeholder', 'image')} ]", first=True,
              align=PP_ALIGN.CENTER, size=18, italic=True, color=GRAY)
    if data.get("caption"):
        tf = _textbox(slide, Inches(0.7), Inches(6.6), Inches(11.9), Inches(0.4))
        _line(tf, data["caption"], first=True, align=PP_ALIGN.CENTER, size=14,
              italic=True, color=GRAY)
    _notes(slide, data.get("notes"))


def render_code(prs, data):
    slide = _blank(prs)
    _title(slide, data["title"])

    lines = data["code"].split("\n")
    has_kicker = bool(data.get("kicker"))
    max_h = 4.35 if has_kicker else 5.05
    card_h = Inches(min(max_h, 0.5 + 0.255 * len(lines)))
    card = _rect(slide, Inches(0.7), Inches(1.62), Inches(11.9), card_h, CARD)
    _gradient(slide, Inches(0.7), Inches(1.62), Inches(0.07), card_h, angle=90.0)

    tf = card.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.35)
    tf.margin_top = Inches(0.22)
    tf.margin_right = tf.margin_bottom = Inches(0.2)
    for i, line in enumerate(lines):
        stripped = line.strip()
        comment = stripped.startswith("#") or stripped.startswith("//")
        _line(tf, line or " ", first=(i == 0), size=14, font=CODE_FONT,
              color=CYAN if comment else CODE_FG)

    if has_kicker:
        tf2 = _textbox(slide, Inches(0.7), Inches(1.62) + card_h + Inches(0.22),
                       Inches(11.9), Inches(0.5))
        _line(tf2, data["kicker"], first=True, size=17, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


def render_demo(prs, data):
    slide = _blank(prs)
    _dark_background(slide)
    _gradient(slide, 0, 0, SLIDE_W, Inches(0.12))
    _picture(slide, "bee.png", left=Inches(10.2), top=Inches(2.3), height=Inches(3.6))

    tf = _textbox(slide, Inches(0.8), Inches(1.1), Inches(8.5), Inches(1.5))
    _line(tf, "DÉMO", first=True, size=88, bold=True, color=WHITE, font=HEADING_FONT)
    _gradient(slide, Inches(0.85), Inches(2.55), Inches(3.2), Inches(0.12))

    tf2 = _textbox(slide, Inches(0.85), Inches(2.95), Inches(8.4), Inches(0.8))
    _line(tf2, data["title"], first=True, size=26, color=CYAN)

    y = Inches(3.9)
    for i, step in enumerate(data.get("steps", [])):
        _rect(slide, Inches(0.88), y + Inches(0.1), Inches(0.17), Inches(0.17),
              CYAN if i % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
        tf3 = _textbox(slide, Inches(1.3), y, Inches(8.4), Inches(0.5))
        _line(tf3, step, first=True, size=19, color=MIST)
        y += Inches(0.62)
    _notes(slide, data.get("notes"))


def render_takeaways(prs, data):
    slide = _blank(prs)
    _honeycomb(slide, Inches(11.45), Inches(1.9))
    _title(slide, data["title"])

    y = Inches(2.0)
    for i, (headline, sub) in enumerate(data["items"]):
        badge = _rect(slide, Inches(0.72), y, Inches(0.62), Inches(0.62),
                      CYAN if i % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
        btf = badge.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _line(btf, str(i + 1), first=True, align=PP_ALIGN.CENTER, size=22,
              bold=True, color=WHITE, font=HEADING_FONT)

        tf = _textbox(slide, Inches(1.65), y - Inches(0.02), Inches(10.8), Inches(0.6))
        _line(tf, headline, first=True, size=27, bold=True, color=INK,
              font=HEADING_FONT)
        tf2 = _textbox(slide, Inches(1.65), y + Inches(0.5), Inches(10.8), Inches(0.6))
        _line(tf2, sub, first=True, size=18, color=GRAY)
        y += Inches(1.45)
    _notes(slide, data.get("notes"))


def _ring_photo(slide, x, y, d, photo, initials):
    """Round photo inside a cyan → purple ring; initials if no photo yet."""
    ring = Inches(0.09)
    _gradient(slide, x - ring, y - ring, d + 2 * ring, d + 2 * ring,
              angle=45.0, shape=MSO_SHAPE.OVAL)
    path = ASSETS / photo if photo else None
    if path and path.exists():
        pic = slide.shapes.add_picture(str(path), x, y, width=d, height=d)
        pic.auto_shape_type = MSO_SHAPE.OVAL
        return
    disc = _rect(slide, x, y, d, d, CARD, shape=MSO_SHAPE.OVAL)
    tf = disc.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _line(tf, initials, first=True, align=PP_ALIGN.CENTER, size=54, bold=True,
          color=WHITE, font=HEADING_FONT)


def render_speakers(prs, data):
    slide = _blank(prs)
    _title(slide, data["title"])
    card_w, card_h, top = Inches(5.75), Inches(4.75), Inches(1.75)
    for i, sp in enumerate(data["speakers"]):
        x = Inches(0.7) + i * (card_w + Inches(0.4))
        _rect(slide, x, top, card_w, card_h, HEX_TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE
              ).adjustments[0] = 0.06
        _gradient(slide, x, top, card_w, Inches(0.09))
        d = Inches(1.9)
        _ring_photo(slide, x + Inches(0.4), top + Inches(0.45), d,
                    sp.get("photo"), sp.get("initials", "?"))
        tx = x + Inches(2.65)
        tf = _textbox(slide, tx, top + Inches(0.55), card_w - Inches(2.85), Inches(1.4))
        first, *rest = sp["name"].split(" ", 1)
        _line(tf, first, first=True, size=30, bold=True, color=INK, font=HEADING_FONT)
        if rest:
            _line(tf, rest[0], size=30, bold=True, color=INK, font=HEADING_FONT)
        ry = top + Inches(2.7)
        for j, role in enumerate(sp["roles"]):
            _rect(slide, x + Inches(0.45), ry + Inches(0.1), Inches(0.16), Inches(0.16),
                  CYAN if j % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
            tf2 = _textbox(slide, x + Inches(0.8), ry, card_w - Inches(1.1), Inches(0.4))
            _line(tf2, role, first=True, size=18, color=BODY_INK, bold=(j == 0))
            ry += Inches(0.45)
        bx, by = tx, top + Inches(1.65)          # badges under the name
        for badge in sp.get("badges", []):
            name, h = badge if isinstance(badge, tuple) else (badge, 0.5)
            pic = _picture(slide, name, left=bx,
                           top=by + (Inches(0.7) - Inches(h)) // 2, height=Inches(h))
            if pic is not None:
                bx += pic.width + Inches(0.35)
        if sp.get("sticker"):                    # tilted on the card's top-right corner
            name, h, rot = sp["sticker"]
            pic = _picture(slide, name, left=x + card_w - Inches(h * 0.75),
                           top=top - Inches(h * 0.3), height=Inches(h))
            if pic is not None:
                pic.rotation = rot
    _notes(slide, data.get("notes"))


def render_journey(prs, data):
    """The story so far: hexagon steps on a line, the bee flies over the current one."""
    slide = _blank(prs)
    _title(slide, data["title"])
    steps, cur = data["steps"], data.get("current", 0)
    n = len(steps)
    left, right, cy = Inches(1.55), Inches(11.75), Inches(3.55)
    gap = (right - left) / (n - 1)
    _rect(slide, left, cy - Inches(0.025), right - left, Inches(0.05), MIST)
    if cur > 0:
        _gradient(slide, left, cy - Inches(0.03), int(gap * cur), Inches(0.06))
    for i, (label, sub) in enumerate(steps):
        cx = left + int(gap * i)
        big = i == cur
        d = Inches(1.05) if big else Inches(0.78)
        if i < cur:
            h = _rect(slide, cx - d / 2, cy - d / 2, d, d,
                      CYAN if i % 2 == 0 else PURPLE, shape=MSO_SHAPE.HEXAGON)
        elif big:
            h = _gradient(slide, cx - d / 2, cy - d / 2, d, d, angle=45.0,
                          shape=MSO_SHAPE.HEXAGON)
        else:
            h = _rect(slide, cx - d / 2, cy - d / 2, d, d, HEX_TINT,
                      shape=MSO_SHAPE.HEXAGON)
            h.line.color.rgb = MIST
            h.line.width = Pt(1.5)
        tf = h.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _line(tf, str(i + 1), first=True, align=PP_ALIGN.CENTER,
              size=26 if big else 20, bold=True,
              color=GRAY if i > cur else WHITE, font=HEADING_FONT)
        w = Inches(2.5)
        tf2 = _textbox(slide, cx - w / 2, cy + Inches(0.75), w, Inches(1.4))
        _line(tf2, label, first=True, align=PP_ALIGN.CENTER, size=17 if big else 15,
              bold=True, color=INK if i <= cur else GRAY, font=HEADING_FONT,
              space_after=4)
        if sub:
            _line(tf2, sub, align=PP_ALIGN.CENTER, size=13,
                  color=BODY_INK if big else GRAY)
        if big:
            _picture(slide, data.get("mascot", "bee_happy.png"),
                     left=cx - Inches(0.45), top=cy - Inches(2.05), height=Inches(1.3))
    if data.get("kicker"):
        _gradient(slide, Inches(0.7), Inches(6.2), Inches(11.9), Inches(0.05))
        tf = _textbox(slide, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5))
        _line(tf, data["kicker"], first=True, size=18, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


def render_compare(prs, data):
    """Two columns: the old way (muted) vs the new way (gradient header)."""
    slide = _blank(prs)
    _title(slide, data["title"])
    col_w, top = Inches(5.75), Inches(1.75)
    col_h = Inches(4.15) if data.get("kicker") else Inches(4.7)
    for c, side in enumerate(("left", "right")):
        col = data[side]
        x = Inches(0.7) + c * (col_w + Inches(0.4))
        # muted_left (default): the left column is "the old way", greyed out.
        # False: two peers — the left header goes solid purple instead.
        muted = c == 0 and data.get("muted_left", True)
        _rect(slide, x, top, col_w, col_h,
              RGBColor(0xF3, 0xF4, 0xF7) if muted else HEX_TINT)
        head_h = Inches(0.7)
        head = (_gradient(slide, x, top, col_w, head_h) if c
                else _rect(slide, x, top, col_w, head_h, GRAY if muted else PURPLE))
        tf = head.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.3)
        _line(tf, col["title"], first=True, size=22, bold=True, color=WHITE,
              font=HEADING_FONT)
        y = top + head_h + Inches(0.3)
        for i, item in enumerate(col["items"]):
            _rect(slide, x + Inches(0.3), y + Inches(0.1), Inches(0.17), Inches(0.17),
                  MIST if muted else (PURPLE if i % 2 else CYAN), shape=MSO_SHAPE.HEXAGON)
            tf = _textbox(slide, x + Inches(0.65), y, col_w - Inches(0.9), Inches(0.6))
            _line(tf, item, first=True, size=18, color=GRAY if muted else BODY_INK)
            y += Inches(0.62) + Inches(0.3) * (len(item) // 44)
    if data.get("kicker"):
        _gradient(slide, Inches(0.7), Inches(6.2), Inches(11.9), Inches(0.05))
        tf = _textbox(slide, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.5))
        _line(tf, data["kicker"], first=True, size=18, bold=True, color=PURPLE)
    _notes(slide, data.get("notes"))


RENDERERS = {
    "title": render_title,
    "speakers": render_speakers,
    "journey": render_journey,
    "compare": render_compare,
    "bullets": render_bullets,
    "image": render_image,
    "code": render_code,
    "demo": render_demo,
    "takeaways": render_takeaways,
    "end": render_end,
}

DARK_TYPES = {"title", "demo", "end"}


# ── motion ───────────────────────────────────────────────────────────────────
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS = (f'xmlns:p="{_P}" '
       'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
       'xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"')

# Gallery is a PowerPoint 2010 transition: readers without p14 fall back to Fade.
_TRANSITION = f"""<mc:AlternateContent {_NS}>
  <mc:Choice Requires="p14">
    <p:transition spd="med" p14:dur="1000"><p14:gallery dir="l"/></p:transition>
  </mc:Choice>
  <mc:Fallback><p:transition spd="med"><p:fade/></p:transition></mc:Fallback>
</mc:AlternateContent>"""


def _zoom_in(spid: int, ids) -> str:
    """Zoom entrance (preset 53, from the object's center), With Previous."""
    def behavior(attr):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr>'
                f'<p:cTn id="{next(ids)}" dur="500" fill="hold"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
                f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst>'
                f'</p:cBhvr><p:tavLst>'
                f'<p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="#{attr}"/></p:val></p:tav>'
                f'</p:tavLst></p:anim>')
    return (f'<p:par><p:cTn id="{next(ids)}" presetID="53" presetClass="entr" '
            f'presetSubtype="16" fill="hold" nodeType="withEffect">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
            f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
            f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            + behavior("ppt_w") + behavior("ppt_h")
            + f'<p:animEffect transition="in" filter="fade"><p:cBhvr>'
            f'<p:cTn id="{next(ids)}" dur="500"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f'</p:childTnLst></p:cTn></p:par>')


def _add_transition(slide):
    sld = slide._element
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(etree.fromstring(_TRANSITION))


def _animate_bees(slide):
    bees = [sh.shape_id for sh in slide.shapes if sh.name.startswith("bee")]
    if not bees:
        return
    ids = iter(range(5, 10_000))
    effects = "".join(_zoom_in(spid, ids) for spid in bees)
    timing = etree.fromstring(
        f'<p:timing {_NS}><p:tnLst><p:par>'
        f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek">'
        f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst>'
        f'<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
        f'</p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst>{effects}</p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn></p:par>'
        f'</p:childTnLst></p:cTn>'
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    # p:timing follows p:transition (or its AlternateContent) in the schema order.
    sld = slide._element
    after = sld.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent")
    if after is None:
        after = sld.find(qn("p:clrMapOvr"))
    after.addnext(timing)


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    total = len(SLIDES)
    cues = getattr(slides_content, "CUES", {})
    unknown = set(cues) - {data.get("title") for data in SLIDES}
    if unknown:
        raise ValueError(f"CUES for unknown slide titles: {sorted(unknown)}")
    for i, data in enumerate(SLIDES):
        renderer = RENDERERS.get(data["type"])
        if renderer is None:
            raise ValueError(f"Slide {i + 1}: unknown type {data['type']!r}")
        renderer(prs, data)
        if cues.get(data.get("title")):
            _cue(prs.slides[i], cues[data["title"]])
        _bees(prs.slides[i], data)
        _footer(prs.slides[i], i, total - 1, dark=data["type"] in DARK_TYPES)
        if getattr(slides_content, "TRANSITION", True):
            _add_transition(prs.slides[i])
        if getattr(slides_content, "BEE_ANIMATION", True):
            _animate_bees(prs.slides[i])
    return prs


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(OUTPUT)
    print(f"OK — {len(prs.slides)} slides → {OUTPUT}")


if __name__ == "__main__":
    main()
