"""Approximate PNG preview of the generated deck — a layout check, not a render.

For when PowerPoint/Keynote cannot export (locked screen, no display). Draws the
shapes, gradients, pictures and word-wrapped text with Pillow and macOS fonts.

Usage (from the repo root):
    uv run --with python-pptx --with pillow python talks/preview_slides.py \
        talks/slides/<DECK_NAME>.pptx /tmp/preview
"""
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC, OUT = sys.argv[1], Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
SCALE = 1100 / 12192000  # px per EMU
AV = "/System/Library/Fonts/Avenir Next.ttc"
MENLO = "/System/Library/Fonts/Menlo.ttc"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_fc = {}


def font(size_pt, bold, code):
    key = (size_pt, bold, code)
    if key not in _fc:
        size = max(6, int(size_pt * 12700 * SCALE))
        path, idx = (MENLO, 0) if code else (AV, 0 if bold else 7)
        _fc[key] = ImageFont.truetype(path, size, index=idx)
    return _fc[key]


def px(v):
    return int(v * SCALE)


def color_of(fill):
    try:
        if fill.type == 1:
            return "#" + str(fill.fore_color.rgb)
        if fill.type == 3:
            st = fill.gradient_stops
            return ("grad", "#" + str(st[0].color.rgb), "#" + str(st[1].color.rgb))
    except Exception:
        pass
    return None


def hexagon(w, h):
    a = w * 0.25
    return [(a, 0), (w - a, 0), (w, h / 2), (w - a, h), (a, h), (0, h / 2)]


def draw_shape(img, sh):
    x, y, w, h = px(sh.left), px(sh.top), max(1, px(sh.width)), max(1, px(sh.height))
    prst = sh._element.spPr.find(A + "prstGeom")
    geom = prst.get("prst") if prst is not None else "rect"
    mask = Image.new("L", (w, h), 0)
    md = ImageDraw.Draw(mask)
    if geom == "hexagon":
        md.polygon(hexagon(w - 1, h - 1), fill=255)
    elif geom == "ellipse":
        md.ellipse([0, 0, w - 1, h - 1], fill=255)
    elif geom == "roundRect":
        md.rounded_rectangle([0, 0, w - 1, h - 1], radius=int(min(w, h) * 0.06), fill=255)
    else:
        md.rectangle([0, 0, w, h], fill=255)
    c = color_of(sh.fill)
    if c:
        if isinstance(c, tuple):
            g = Image.linear_gradient("L")
            g = g.resize((w, h)) if h > w * 1.5 else g.rotate(90).resize((w, h))
            layer = Image.composite(Image.new("RGB", (w, h), c[2]),
                                    Image.new("RGB", (w, h), c[1]), g)
        else:
            layer = Image.new("RGB", (w, h), c)
        img.paste(layer, (x, y), mask)
    try:
        if sh.line.fill.type == 1:
            ImageDraw.Draw(img).polygon([(x + a, y + b) for a, b in hexagon(w - 1, h - 1)],
                                        outline="#" + str(sh.line.color.rgb))
    except Exception:
        pass


def draw_text(img, sh):
    if not sh.has_text_frame:
        return
    tf = sh.text_frame
    x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
    ml, mt, mr = px(tf.margin_left or 0), px(tf.margin_top or 0), px(tf.margin_right or 0)
    d = ImageDraw.Draw(img)
    maxw = (w - ml - mr) if tf.word_wrap is not False else 10**6
    lines = []
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            r0 = p.runs[0] if p.runs else None
            lines.append(([], r0.font.size.pt if r0 is not None and r0.font.size else 18, p, 0))
            continue
        size = max(r.font.size.pt if r.font.size else 18 for r in runs)
        cur, cur_w = [], 0
        for r in runs:
            f = font(r.font.size.pt if r.font.size else 18, bool(r.font.bold),
                     r.font.name == "Consolas")
            col = "#" + str(r.font.color.rgb) if r.font.color and r.font.color.type else "#000"
            for word in r.text.split(" "):
                token = word + " "
                tw = d.textlength(token, font=f)
                if cur and cur_w + tw > maxw:
                    lines.append((cur, size, p, cur_w))
                    cur, cur_w = [], 0
                cur.append((token, f, col))
                cur_w += tw
        lines.append((cur, size, p, cur_w))
    step = lambda s, p: px(s * 12700 * 1.2) + px(p.space_after or 0)  # noqa: E731
    total = sum(step(s, p) for _, s, p, _ in lines)
    top = y + mt
    if tf.vertical_anchor == 3:
        top = y + (h - total) / 2
    elif tf.vertical_anchor == 4:
        top = y + h - total
    for toks, s, p, lw in lines:
        lx = x + ml
        if p.alignment == 2:
            lx = x + (w - lw) / 2
        elif p.alignment == 3:
            lx = x + w - lw
        for t, f, col in toks:
            d.text((lx, top), t, font=f, fill=col)
            lx += d.textlength(t, font=f)
        top += step(s, p)


def draw_pic(img, sh):
    w, h = max(1, px(sh.width)), max(1, px(sh.height))
    im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA").resize((w, h))
    xfrm = sh._element.spPr.find(A + "xfrm")
    if xfrm is not None and xfrm.get("flipH") == "1":
        im = im.transpose(Image.FLIP_LEFT_RIGHT)
    prst = sh._element.spPr.find(A + "prstGeom")
    if prst is not None and prst.get("prst") == "ellipse":
        m = Image.new("L", (w, h), 0)
        ImageDraw.Draw(m).ellipse([0, 0, w - 1, h - 1], fill=255)
        im.putalpha(Image.composite(im.getchannel("A"), m, m))
    cx, cy = px(sh.left) + w / 2, px(sh.top) + h / 2
    if sh.rotation:
        im = im.rotate(-sh.rotation, expand=True, resample=Image.BICUBIC)
    img.paste(im, (int(cx - im.width / 2), int(cy - im.height / 2)), im)


prs = Presentation(SRC)
for i, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (px(prs.slide_width), px(prs.slide_height)), "white")
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            draw_pic(img, sh)
        else:
            draw_shape(img, sh)
            draw_text(img, sh)
    img.save(OUT / f"s{i:02d}.png")
print(f"{len(prs.slides)} previews → {OUT}")
